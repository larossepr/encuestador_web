from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
import os

def crear_parrafo(texto):
    p = doc.add_paragraph(texto)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.runs[0] if p.runs else p.add_run(texto)
    run.font.name = 'Arial'
    run.font.size = Pt(14)
    return p
def formato_tabla(tabla):
    for row in tabla.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    run.font.name = 'Arial'
                    run.font.size = Pt(14)
doc = Document()
nombre = input("Nombre del proyecto= ").title()

titulo = doc.add_paragraph(f"Análisis de Interesados de las Instalaciones del Proyecto {nombre}")  
titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
run_titulo = titulo.runs[0] if titulo.runs else titulo.add_run(f"Análisis de Interesados de las Instalaciones del Proyecto: {nombre}")
run_titulo.font.size = Pt(18)
run_titulo.font.name = 'Arial'
run_titulo.underline = True
run_titulo.bold = True
p1 = crear_parrafo("Sexo: Femenino____ Masculino____")
p2 = crear_parrafo("Edad: ____________")
p3 = crear_parrafo("Estado Civil: Soltero____ Casado____ Viudo____ Divorciado____ Union Libre____")
p4 = crear_parrafo("Nivel de estudios: Basica____ Secundaria____ Bachiller____ Superior____")
p5 = crear_parrafo("Profesion_______________ ")
p6 = crear_parrafo("En que sector trabaja: Publico____ Privado____ Independiente____ Desempleado____")
p7 = crear_parrafo("Lugares que la gente visita en la comunidad para divertirse: Colmados____ Bares____ Galleras____ Play____ Billar____ Playa____ Rio____ Otros: ____________")
p8 = crear_parrafo("Pertenece usted a alguna organizacion, seleccione entre las siguientes: Iglesia____ Junta de Vecinos____ Club Deportivo____ Asociacion de Agricultores____ Asociacion de Comerciantes____ Asociacion de Ganaderos____ Otro: ____________")
p9 = crear_parrafo("Posee usted tierra? Si____ No____ Posee usted casa? Si____ No____")
p10 = crear_parrafo(f"Sabia usted de la existencia del proyecto {nombre} Si____ No____")
p11 = crear_parrafo(f"Cree usted que el proyecto {nombre} es positivo para la comunidad? Si____ No____")
p12 = crear_parrafo("Como considera usted que es la calidad actual de estos componentes fudamentales: ")
tabla = doc.add_table(rows=5, cols=4 )
tabla.cell(0, 0).text = ""
tabla.cell(0, 1).text = "Buena"
tabla.cell(0, 2).text = "Regular"
tabla.cell(0, 3).text = "Mala"
tabla.cell(1, 0).text = "Aire"
tabla.cell(1, 1).text = "____"
tabla.cell(1, 2).text = "____"
tabla.cell(1, 3).text = "____"
tabla.cell(2, 0).text = "Agua"
tabla.cell(2, 1).text = "____"
tabla.cell(2, 2).text = "____"
tabla.cell(2, 3).text = "____"
tabla.cell(3, 0).text = "Tierra"
tabla.cell(3, 1).text = "____"
tabla.cell(3, 2).text = "____"
tabla.cell(3, 3).text = "____"
tabla.cell(4, 0).text = "Paisaje"
tabla.cell(4, 1).text = "____"
tabla.cell(4, 2).text = "____"
tabla.cell(4, 3).text = "____"
formato_tabla(tabla)
p13 = crear_parrafo(f"Cree usted que las actividades del proyecto {nombre} tiene incidencia sobre el paisaje? Lo Mejorara____ Lo Dañara____ No Incide____")
p14 = crear_parrafo(f"Cree usted que las actividades del proyecto {nombre} afecta a los animales? Si____ No____")
p15 = crear_parrafo(f"Cree usted que las actividades del proyecto {nombre} afecta o afectara a los arboles y a las plantas? Si____ No____")
p16 = crear_parrafo("Cree usted que esta zona podria inundarse por algun rio o tsunami del mar? Si____ No____")
p17 = crear_parrafo(f"Cree usted que la instalacion y las actividades del proyecto {nombre} contribuye al desarollo de la comunidad? Si____ No____")
p18 = crear_parrafo(f"Cree usted que las actividades del proyecto {nombre}, ayuda a la creacion de empleo? Si____ No____")
p19 = crear_parrafo("Es usted nativo(a) de aqui? Si____ No____")
p20 = crear_parrafo("Reside en esta Comunidad? Si____ No____ (En caso de que no resida): Viene por temporada? Si____ No____")
p21 = crear_parrafo("Podria decirnos su nombre y direccion?")
p21_1 = crear_parrafo("Nombre: __________________________________________________")
p21_2 = crear_parrafo("Direccion: ________________________________________________________")

output_dir = "encuestas"
os.makedirs(output_dir, exist_ok=True)
encuesta = os.path.join(output_dir, f"{nombre}_encuesta.docx")

doc.save(encuesta)

